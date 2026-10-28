"""Демо-материалы для веб-демо Deixis: лекция (PPTX → PDF через LibreOffice → страницы JPEG), конспект Word
(страницы — свой перевод DOCX из ядра), фото и видео.

Всё рисуется здесь же, без чужих картинок: python-pptx, Pillow, numpy, LibreOffice, pdftoppm, ffmpeg.
Запуск: python3 tools/demo/make_assets.py  (результат — tools/demo/media/)
"""
import math
import os
import shutil
import subprocess
import sys
import tempfile

import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "media")
FONT = os.path.join(HERE, "..", "..", "core", "src", "main", "resources", "web", "NotoSans-SemiBold.ttf")

INK = RGBColor(0x1B, 0x23, 0x2E)
MUTED = RGBColor(0x5B, 0x66, 0x73)
ACCENT = RGBColor(0x1F, 0x6F, 0x8B)
WARM = RGBColor(0xD9, 0x82, 0x3B)
PAPER = RGBColor(0xFA, 0xF8, 0xF4)


def lecture(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def slide():
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = PAPER
        return s

    def text(s, x, y, w, h, lines, size, color=INK, bold=False, align=PP_ALIGN.LEFT, space=6):
        tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space)
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Noto Sans"
        return tb

    def bar(s):
        r = s.shapes.add_shape(1, 0, 0, Emu(W), Emu(180000))
        r.fill.solid()
        r.fill.fore_color.rgb = ACCENT
        r.line.fill.background()

    m = 760000
    # 1. Титул
    s = slide()
    bar(s)
    text(s, m, 1900000, W - 2 * m, 600000, ["Урбанистика · лекция 3"], 22, MUTED)
    text(s, m, 2450000, W - 2 * m, 1500000, ["Город пятнадцати минут"], 54, INK, bold=True)
    text(s, m, 3900000, W - 2 * m, 900000, ["Как расстояние до школы, врача и магазина меняет жизнь района"], 24, MUTED)

    # 2. Тезисы
    s = slide()
    bar(s)
    text(s, m, 600000, W - 2 * m, 900000, ["Три вопроса жителя"], 40, INK, bold=True)
    text(s, m, 1750000, W - 2 * m, 3800000, [
        "1. Можно ли дойти до всего нужного пешком?",
        "2. Безопасно ли идти — тротуары, свет, переходы?",
        "3. Есть ли куда выйти — парк, площадь, двор?",
    ], 30, INK, space=18)

    # 3. Диаграмма
    s = slide()
    bar(s)
    text(s, m, 500000, W - 2 * m, 900000, ["Сколько поездок делают пешком"], 40, INK, bold=True)
    data = CategoryChartData()
    data.categories = ["Окраина", "Спальный район", "Центр", "Квартал 15 минут"]
    data.add_series("Доля поездок пешком, %", (12, 21, 38, 54))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(m), Emu(1500000), Emu(W - 2 * m), Emu(4700000), data)
    ch = gf.chart
    ch.has_legend = False
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.font.size = Pt(20)
    ch.plots[0].data_labels.font.color.rgb = INK
    ch.plots[0].gap_width = 60
    ch.series[0].format.fill.solid()
    ch.series[0].format.fill.fore_color.rgb = ACCENT
    pt = ch.series[0].points[3]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = WARM
    ch.category_axis.tick_labels.font.size = Pt(18)
    ch.category_axis.tick_labels.font.color.rgb = MUTED
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    text(s, m, 6250000, W - 2 * m, 400000, ["Условные данные для примера"], 14, MUTED)

    # 4. Таблица
    s = slide()
    bar(s)
    text(s, m, 500000, W - 2 * m, 900000, ["Что должно быть рядом"], 40, INK, bold=True)
    rows = [("Место", "Пешком", "Как часто"), ("Магазин", "5 минут", "каждый день"), ("Школа", "10 минут", "каждый день"),
            ("Поликлиника", "15 минут", "раз в месяц"), ("Парк", "10 минут", "каждую неделю")]
    tbl = s.shapes.add_table(len(rows), 3, Emu(m), Emu(1600000), Emu(W - 2 * m), Emu(4200000)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = ACCENT if i == 0 else (RGBColor(0xE6, 0xEF, 0xF2) if i % 2 else RGBColor(0xF4, 0xF7, 0xF8))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(24)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK

    # 5. Итог
    s = slide()
    bar(s)
    text(s, m, 2000000, W - 2 * m, 1200000, ["Город удобен, когда до нужного — шаги, а не пробки"], 44, INK, bold=True)
    text(s, m, 4000000, W - 2 * m, 600000, ["Вопросы и обсуждение"], 26, ACCENT)
    prs.save(path)


SUNSET = ((28, 42, 78), (196, 110, 92), (246, 196, 120), [(118, 84, 104), (84, 60, 86), (52, 40, 64), (26, 22, 38)], 7)
MORNING = ((120, 170, 210), (190, 214, 226), (236, 232, 214), [(110, 140, 150), (78, 110, 118), (52, 82, 88), (30, 52, 56)], 11)


def report(path):
    import docx
    from docx.shared import Pt as DPt
    d = docx.Document()
    d.add_heading("Отчёт по практике", 0)
    d.add_paragraph("Студент 3 курса, направление «Градостроительство». Тема: доступность школ и поликлиник в районе.")
    d.add_heading("Что сделано", 1)
    for line in ("Обследованы 12 кварталов и 46 маршрутов пешком.", "Замерено время до школы, поликлиники и парка.", "Составлена карта мест, куда трудно дойти пешком."):
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("Итог", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Light Grid Accent 1"
    for r, row in enumerate((("Место", "Среднее время", "Норма"), ("Школа", "11 минут", "до 10"), ("Поликлиника", "19 минут", "до 15"), ("Парк", "8 минут", "до 10"))):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    p = d.add_paragraph("Условные данные для примера.")
    p.runs[0].font.size = DPt(9)
    d.save(path)


def photo(path, palette=SUNSET, w=1920, h=1080):
    """Горы: несколько слоёв силуэтов над небом с переходом цвета."""
    y = np.linspace(0, 1, h)[:, None]
    top, mid, low = (np.array(c) / 255 for c in palette[:3])
    t = np.clip(y / 0.62, 0, 1)
    sky = np.where(t < 0.6, top + (mid - top) * (t / 0.6), mid + (low - mid) * ((t - 0.6) / 0.4))
    img = np.broadcast_to(sky[:, None, :], (h, w, 3)).copy()
    # Солнце
    yy, xx = np.mgrid[0:h, 0:w]
    d = np.hypot(xx - w * 0.68, yy - h * 0.6)
    glow = np.exp(-(d / (h * 0.22)) ** 2)[..., None]
    img = img + glow * (np.array([1.0, 0.86, 0.6]) - img) * 0.85
    img = np.clip(img, 0, 1)
    pil = Image.fromarray((img * 255).astype(np.uint8))
    draw = ImageDraw.Draw(pil)
    rng = np.random.default_rng(palette[4])
    layers = list(zip((0.55, 0.66, 0.76, 0.88), palette[3], (90, 70, 55, 40)))
    for base, color, amp in layers:
        pts = []
        phase = rng.uniform(0, 6.28, 4)
        for x in range(0, w + 8, 8):
            v = sum(math.sin(x / (w / (k + 1.3)) + phase[k]) * amp / (k + 1) for k in range(4))
            v += rng.normal(0, amp * 0.04)
            pts.append((x, h * base - v))
        draw.polygon(pts + [(w, h), (0, h)], fill=color)
    pil = pil.filter(ImageFilter.GaussianBlur(0.6))
    pil.save(path, quality=88)


def video(path, font):
    """12 секунд: плавный переход цвета и подпись. H.264, без звука, открывается в любом браузере."""
    vf = (f"drawtext=fontfile={font}:text='Deixis':fontcolor=white:fontsize=120:x=(w-text_w)/2:y=(h-text_h)/2-50,"
          f"drawtext=fontfile={font}:text='Демо-видео · %{{eif\\:t\\:d}} с':fontcolor=white@0.8:fontsize=44:x=(w-text_w)/2:y=(h/2)+70")
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                    "gradients=s=1280x720:c0=1f6f8b:c1=d9823b:c2=1b232e:x0=0:y0=0:x1=1280:y1=720:nb_colors=3:speed=0.02:d=12:r=30",
                    "-vf", vf, "-c:v", "libx264", "-profile:v", "main", "-pix_fmt", "yuv420p", "-crf", "30",
                    "-movflags", "+faststart", path], check=True)


def konspekt(tmp):
    """Демо-документ Word с телефона (make_docx.py): страницы рисует свой перевод DOCX из ядра, как на телефоне."""
    root = os.path.abspath(os.path.join(HERE, "..", ".."))
    src = os.path.join(root, "android", "phone", "src", "main", "assets", "demo", "Конспект. Город 15 минут.docx")
    out = os.path.join(tmp, "konspekt")
    env = dict(os.environ, LC_ALL="C.UTF-8", DEIXIS_SKIP_ANDROID="1")
    subprocess.run([os.path.join(root, "gradlew"), "-q", ":core:test", "--tests", "*DocxTest.rendersPages", "--rerun",
                    f"-Ddeixis.docxPng={out}", f"-Ddeixis.docxFile={src}"], cwd=root, env=env, check=True)
    pages = sorted(f for f in os.listdir(out) if f.startswith("docx-page-"))
    for i, f in enumerate(pages):
        Image.open(os.path.join(out, f)).convert("RGB").save(os.path.join(OUT, f"konspekt-{i}.jpg"), quality=85)


def main():
    os.makedirs(OUT, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pptx = os.path.join(tmp, "lecture.pptx")
        lecture(pptx)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            sys.exit("Нужен LibreOffice, чтобы перевести лекцию в PDF")
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "pdf", "--outdir", tmp, pptx],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        shutil.copy(os.path.join(tmp, "lecture.pdf"), os.path.join(OUT, "lecture.pdf"))
        shutil.copy(pptx, os.path.join(OUT, "lecture.pptx"))
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=86", "-scale-to-x", "1920", "-scale-to-y", "-1",
                        os.path.join(tmp, "lecture.pdf"), os.path.join(tmp, "p")], check=True)
        pages = sorted(f for f in os.listdir(tmp) if f.startswith("p-") and f.endswith(".jpg"))
        for i, f in enumerate(pages):
            shutil.copy(os.path.join(tmp, f), os.path.join(OUT, f"lecture-{i}.jpg"))
        print("страниц:", len(pages))
        # Документ для перевода на компьютере: одна страница отчёта.
        docx = os.path.join(tmp, "report.docx")
        report(docx)
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "pdf", "--outdir", tmp, docx],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=86", "-scale-to-x", "1240", "-scale-to-y", "-1", "-f", "1", "-l", "1",
                        os.path.join(tmp, "report.pdf"), os.path.join(tmp, "r")], check=True)
        shutil.copy(os.path.join(tmp, sorted(f for f in os.listdir(tmp) if f.startswith("r-"))[0]), os.path.join(OUT, "report-0.jpg"))
        konspekt(tmp)
    photo(os.path.join(OUT, "mountains.jpg"))
    photo(os.path.join(OUT, "morning.jpg"), MORNING)
    video(os.path.join(OUT, "clip.mp4"), os.path.abspath(FONT))
    for f in sorted(os.listdir(OUT)):
        print(f, os.path.getsize(os.path.join(OUT, f)))


if __name__ == "__main__":
    main()
